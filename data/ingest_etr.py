"""Parse a user-supplied Establish The Run (ETR) rankings export.

Handles both `.xlsx` (ETR's native cheat-sheet download) and `.csv`. For
workbooks, the default sheet is ETR's "Full Ledger" tab -- the machine-
readable one carrying ADP; the "Cheat Sheet" tab is a printable six-column
layout of the same players, colour-coded by position, with no extra data.

Column matching is alias-based (case/whitespace-insensitive) rather than
hardcoded, and callers can override any column ETR renames.

Note that ETR's Top-300 export carries no projected points and no tier
column -- see `data/projections.py`, which turns these ranks into the point
values VBD needs.

Canonical player IDs are synthesized from name+team+position (a stable,
human-readable slug) rather than pulled from an external ID system, so this
module has no network dependency and the wider engine stays testable
offline. `data/id_mapping.py` is responsible for later linking these IDs to
platform-specific IDs (e.g. Sleeper) for live draft sync.
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

from data.models import Player, PlayerRanking, Position

DEFAULT_SHEET_NAME = "Full Ledger"

# Canonical field -> acceptable column header aliases (matched case/whitespace-insensitively).
DEFAULT_COLUMN_ALIASES: dict[str, list[str]] = {
    "name": ["player", "name", "player name"],
    "team": ["team", "nfl team", "tm"],
    "position": ["pos", "position"],
    "rank": ["rank", "overall rank", "ecr", "overall"],
    "tier": ["tier"],
    "projected_points": ["proj", "projection", "points", "fpts", "projected points"],
    "pos_rank": ["pos rank", "positional rank", "position rank"],
    "adp": ["adp", "average draft position"],
}

REQUIRED_FIELDS = ("name", "position", "rank")

_POSITION_ALIASES: dict[str, Position] = {
    "DST": Position.DEF,
    "D/ST": Position.DEF,
    "DEFENSE": Position.DEF,
    **{p.value: p for p in Position},
}


class ETRIngestError(ValueError):
    pass


def _normalize(header: str) -> str:
    return re.sub(r"\s+", " ", header.strip().lower())


def slugify_player_id(name: str, team: str | None, position: Position) -> str:
    slug_name = re.sub(r"[^a-z0-9]+", "-", name.strip().lower()).strip("-")
    parts = [slug_name, (team or "fa").strip().lower(), position.value.lower()]
    return "-".join(parts)


def resolve_column_map(
    fieldnames: list[str], column_overrides: dict[str, str] | None = None
) -> dict[str, str]:
    """Map each canonical field name to the actual CSV column name to use."""
    overrides = column_overrides or {}
    normalized_to_actual = {_normalize(h): h for h in fieldnames}

    resolved: dict[str, str] = {}
    for field, aliases in DEFAULT_COLUMN_ALIASES.items():
        if field in overrides:
            if overrides[field] not in fieldnames:
                raise ETRIngestError(f"column override {overrides[field]!r} for {field!r} not found in CSV headers {fieldnames}")
            resolved[field] = overrides[field]
            continue
        for alias in aliases:
            actual = normalized_to_actual.get(_normalize(alias))
            if actual is not None:
                resolved[field] = actual
                break

    missing_required = [f for f in REQUIRED_FIELDS if f not in resolved]
    if missing_required:
        raise ETRIngestError(
            f"could not find required column(s) {missing_required} in CSV headers {fieldnames}; "
            "pass column_overrides={{'field': 'actual header'}} to fix"
        )
    return resolved


def _parse_position(raw: str) -> Position:
    key = raw.strip().upper()
    if key not in _POSITION_ALIASES:
        valid = sorted(set(_POSITION_ALIASES))
        raise ETRIngestError(f"unrecognized position {raw!r}; expected one of {valid}")
    return _POSITION_ALIASES[key]


def _parse_optional_int(raw: str | None) -> int | None:
    if raw is None or not raw.strip():
        return None
    return int(float(raw))


def _parse_optional_float(raw: str | None) -> float | None:
    if raw is None or not raw.strip():
        return None
    return float(raw)


def _parse_pos_rank(raw: str | None) -> int | None:
    """ETR writes positional rank as e.g. "RB01" / "WR07"; pull out the number."""
    if raw is None or not raw.strip():
        return None
    match = re.search(r"(\d+)", raw)
    return int(match.group(1)) if match else None


def _read_csv_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    # utf-8-sig, not utf-8: ETR's exports carry a BOM, which would otherwise
    # ride along on the first header ("﻿Player") and defeat alias matching
    # for whichever column happens to come first.
    with path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        if not reader.fieldnames:
            raise ETRIngestError(f"{path} has no header row")
        return list(reader.fieldnames), [dict(row) for row in reader]


def _read_xlsx_rows(path: Path, sheet_name: str) -> tuple[list[str], list[dict[str, str]]]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ETRIngestError(
            f"{path} has no sheet named {sheet_name!r}; available sheets: {workbook.sheetnames}"
        )
    sheet = workbook[sheet_name]

    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise ETRIngestError(f"{path} sheet {sheet_name!r} is empty") from None

    fieldnames = ["" if cell is None else str(cell).strip() for cell in header_row]
    rows = []
    for raw_row in rows_iter:
        if all(cell is None for cell in raw_row):
            continue
        rows.append(
            {
                field: ("" if cell is None else str(cell))
                for field, cell in zip(fieldnames, raw_row)
                if field
            }
        )
    workbook.close()
    return [f for f in fieldnames if f], rows


def load_etr_rankings(
    path: str | Path,
    source: str = "ETR",
    column_overrides: dict[str, str] | None = None,
    sheet_name: str = DEFAULT_SHEET_NAME,
) -> tuple[list[Player], list[PlayerRanking]]:
    """Parse an ETR rankings export (.xlsx or .csv) into (players, rankings), aligned by player_id."""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        fieldnames, rows = _read_xlsx_rows(path, sheet_name)
    else:
        fieldnames, rows = _read_csv_rows(path)

    duplicate_headers = {h for h in fieldnames if fieldnames.count(h) > 1}
    if duplicate_headers:
        raise ETRIngestError(f"{path} has duplicate column header(s): {sorted(duplicate_headers)}")
    column_map = resolve_column_map(fieldnames, column_overrides)

    players: list[Player] = []
    rankings: list[PlayerRanking] = []
    seen_player_ids: dict[str, str] = {}
    for row_num, row in enumerate(rows, start=2):
        name = row[column_map["name"]].strip()
        if not name:
            continue
        team = row[column_map["team"]].strip() if "team" in column_map else None
        position = _parse_position(row[column_map["position"]])
        rank_raw = row[column_map["rank"]]
        try:
            rank = int(float(rank_raw))
        except ValueError as exc:
            raise ETRIngestError(f"{path} row {row_num}: invalid rank {rank_raw!r} for {name!r}") from exc

        def _optional(field: str, parse):
            column = column_map.get(field)
            return parse(row.get(column)) if column else None

        tier = _optional("tier", _parse_optional_int)
        projected_points = _optional("projected_points", _parse_optional_float)
        pos_rank = _optional("pos_rank", _parse_pos_rank)
        adp = _optional("adp", _parse_optional_float)

        player_id = slugify_player_id(name, team, position)
        if player_id in seen_player_ids:
            raise ETRIngestError(
                f"{path} row {row_num}: {name!r} ({team}, {position.value}) collides with row "
                f"{seen_player_ids[player_id]} on player_id {player_id!r}"
            )
        seen_player_ids[player_id] = str(row_num)
        players.append(
            Player(
                player_id=player_id,
                name=name,
                nfl_team=team or None,
                position=position,
                eligible_positions=frozenset({position}),
            )
        )
        rankings.append(
            PlayerRanking(
                player_id=player_id,
                source=source,
                rank=rank,
                tier=tier,
                position=position,
                projected_points=projected_points,
                pos_rank=pos_rank,
                adp=adp,
            )
        )

    return players, rankings


# ETR publishes one auction-value column per scoring format. FFL-NY awards 0.5
# per reception (rule 3.4.8), so "half_ppr" is the correct default here; the
# superflex columns never apply, since this league starts a single QB.
AUCTION_VALUE_ALIASES: dict[str, list[str]] = {
    "half_ppr": ["etr half ppr", "half ppr", "half"],
    "full_ppr": ["etr full ppr", "full ppr", "ppr", "full"],
    "std": ["etr std", "std", "standard"],
    "superflex_half": ["etr superflex half", "superflex half"],
    "superflex_full": ["etr superflex full", "superflex full"],
}


def _resolve_column(fieldnames: list[str], aliases: list[str], label: str) -> str:
    normalized_to_actual = {_normalize(h): h for h in fieldnames}
    for alias in aliases:
        if alias in normalized_to_actual:
            return normalized_to_actual[alias]
    raise ETRIngestError(
        f"no column found for {label!r}; tried {aliases} against {fieldnames}"
    )


def load_etr_auction_values(
    path: str | Path,
    value_column: str = "half_ppr",
    sheet_name: str = DEFAULT_SHEET_NAME,
    column_overrides: dict[str, str] | None = None,
) -> tuple[list[Player], dict[str, float]]:
    """Parse an ETR *auction values* export into (players, player_id -> value).

    A separate entry point from `load_etr_rankings` because the auction export
    is a different product with a different shape: it carries a dollar value
    per scoring format and no overall rank at all, so it cannot satisfy the
    rankings loader's required columns.

    The values it returns are a market-clearing allocation of a from-scratch
    auction -- ETR's Half PPR column sums to exactly $2,400, being 12 teams by
    a $200 cap. They are therefore *not* directly usable in a keeper league;
    pass them through `engine.auction.implied_vorp` and re-clear against the
    money and slots the auction will actually have.
    """
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        fieldnames, rows = _read_xlsx_rows(path, sheet_name)
    else:
        fieldnames, rows = _read_csv_rows(path)

    if value_column not in AUCTION_VALUE_ALIASES:
        valid = ", ".join(sorted(AUCTION_VALUE_ALIASES))
        raise ETRIngestError(f"unknown value_column {value_column!r}; expected one of {valid}")

    overrides = column_overrides or {}
    name_col = overrides.get("name") or _resolve_column(fieldnames, DEFAULT_COLUMN_ALIASES["name"], "name")
    pos_col = overrides.get("position") or _resolve_column(fieldnames, DEFAULT_COLUMN_ALIASES["position"], "position")
    team_col = overrides.get("team")
    if team_col is None:
        try:
            team_col = _resolve_column(fieldnames, DEFAULT_COLUMN_ALIASES["team"], "team")
        except ETRIngestError:
            team_col = None
    value_col = overrides.get(value_column) or _resolve_column(
        fieldnames, AUCTION_VALUE_ALIASES[value_column], value_column
    )

    players: list[Player] = []
    values: dict[str, float] = {}
    for row_num, row in enumerate(rows, start=2):
        name = (row.get(name_col) or "").strip()
        if not name:
            continue
        position = _parse_position(row[pos_col])
        team = (row.get(team_col) or "").strip() or None if team_col else None
        value = _parse_optional_float(row.get(value_col))
        if value is None:
            continue

        player_id = slugify_player_id(name, team, position)
        if player_id in values:
            raise ETRIngestError(
                f"{path} row {row_num}: duplicate player_id {player_id!r} for {name!r}"
            )
        players.append(
            Player(
                player_id=player_id,
                name=name,
                nfl_team=team,
                position=position,
                eligible_positions=frozenset({position}),
            )
        )
        values[player_id] = value
    if not values:
        raise ETRIngestError(f"{path} yielded no auction values from column {value_col!r}")
    return players, values
