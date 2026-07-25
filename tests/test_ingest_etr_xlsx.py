"""xlsx ingestion, exercised against a workbook shaped like ETR's real export."""
import openpyxl
import pytest

from data.ingest_etr import ETRIngestError, load_etr_rankings
from data.models import Position

FULL_LEDGER_HEADERS = [
    "Rank", "Player", "Pos", "Team", "Pos Rank", "ADP", "Rank Diff", "ADP Pos Rank", "Pos Rank Diff"
]
FULL_LEDGER_ROWS = [
    [1, "Jahmyr Gibbs", "RB", "DET", "RB01", 1.3, 0.3, "RB01", 0],
    [2, "Puka Nacua", "WR", "LA", "WR01", 4.0, 2.0, "WR02", 1],
    [3, "Brock Bowers", "TE", "LV", "TE01", 20.0, 17.0, "TE01", 0],
    [4, "LAC DST", "DST", "LAC", "DST12", 201.0, 197.0, "DST12", 0],
]


@pytest.fixture
def etr_workbook(tmp_path):
    path = tmp_path / "etr.xlsx"
    wb = openpyxl.Workbook()
    cheat = wb.active
    cheat.title = "Cheat Sheet"
    cheat.append(["ETR Full-PPR Top 300"])

    ledger = wb.create_sheet("Full Ledger")
    ledger.append(FULL_LEDGER_HEADERS)
    for row in FULL_LEDGER_ROWS:
        ledger.append(row)
    wb.save(path)
    return path


def test_reads_full_ledger_sheet_by_default(etr_workbook):
    players, rankings = load_etr_rankings(etr_workbook)
    assert len(players) == 4
    assert [p.name for p in players][:2] == ["Jahmyr Gibbs", "Puka Nacua"]


def test_captures_adp_and_positional_rank(etr_workbook):
    _, rankings = load_etr_rankings(etr_workbook)
    gibbs = rankings[0]
    assert gibbs.pos_rank == 1
    assert gibbs.adp == pytest.approx(1.3)
    assert gibbs.rank_diff == pytest.approx(0.3)


def test_pos_rank_parsed_out_of_padded_label(etr_workbook):
    _, rankings = load_etr_rankings(etr_workbook)
    assert rankings[3].pos_rank == 12  # "DST12"


def test_dst_maps_to_def_position(etr_workbook):
    players, _ = load_etr_rankings(etr_workbook)
    assert players[3].position == Position.DEF


def test_export_without_projections_leaves_points_none(etr_workbook):
    """ETR's Top-300 has no points column; data/projections.py supplies them."""
    _, rankings = load_etr_rankings(etr_workbook)
    assert all(r.projected_points is None for r in rankings)
    assert all(r.tier is None for r in rankings)


def test_rank_diff_is_none_without_adp():
    from data.models import PlayerRanking

    ranking = PlayerRanking(
        player_id="x", source="ETR", rank=5, tier=None, position=Position.RB
    )
    assert ranking.rank_diff is None


def test_missing_sheet_name_raises(etr_workbook):
    with pytest.raises(ETRIngestError, match="no sheet named"):
        load_etr_rankings(etr_workbook, sheet_name="Nope")


def test_blank_rows_are_skipped(tmp_path):
    path = tmp_path / "gaps.xlsx"
    wb = openpyxl.Workbook()
    ledger = wb.active
    ledger.title = "Full Ledger"
    ledger.append(FULL_LEDGER_HEADERS)
    ledger.append(FULL_LEDGER_ROWS[0])
    ledger.append([None] * len(FULL_LEDGER_HEADERS))
    ledger.append(FULL_LEDGER_ROWS[1])
    wb.save(path)

    players, _ = load_etr_rankings(path)
    assert len(players) == 2
